"""Generate a portable DOCX resume from structured sections."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt


@dataclass(frozen=True, slots=True)
class DocxExportResult:
    docx_bytes: bytes
    sha256: str
    size_bytes: int
    extracted_text_hash: str


class StructuredDocxExporter:
    def generate(
        self,
        *,
        title: str,
        subtitle: str,
        sections: list[tuple[str, list[str]]],
    ) -> DocxExportResult:
        document = Document()
        section = document.sections[0]
        section.top_margin = Mm(16)
        section.bottom_margin = Mm(16)
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)

        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(10)

        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(title)
        run.bold = True
        run.font.size = Pt(18)

        if subtitle:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(subtitle)
            run.font.size = Pt(9)

        extracted = [title, subtitle]
        for label, lines in sections:
            if not lines:
                continue
            paragraph = document.add_paragraph()
            run = paragraph.add_run(label)
            run.bold = True
            run.font.size = Pt(12)
            extracted.append(label)
            for line in lines:
                document.add_paragraph(line, style="List Bullet")
                extracted.append(line)

        stream = BytesIO()
        document.save(stream)
        content = stream.getvalue()
        return DocxExportResult(
            docx_bytes=content,
            sha256=hashlib.sha256(content).hexdigest(),
            size_bytes=len(content),
            extracted_text_hash=hashlib.sha256(
                "\n".join(extracted).encode("utf-8")
            ).hexdigest(),
        )

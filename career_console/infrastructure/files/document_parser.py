"""Strict parsers for resume document formats."""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import chardet
from docx import Document as DocxDocument
from pypdf import PdfReader

from career_console.domain.common.errors import CareerDomainError

SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    file_name: str
    media_type: str
    text: str
    parser_name: str


class DocumentParser:
    def __init__(self, *, max_bytes: int) -> None:
        self.max_bytes = max_bytes

    def parse(
        self, *, file_name: str, content: bytes, media_type: str | None = None
    ) -> ParsedDocument:
        safe_name = self._validate_file_name(file_name)
        if not content:
            raise CareerDomainError("The imported document is empty.", code="empty_document")
        if len(content) > self.max_bytes:
            raise CareerDomainError(
                f"Document exceeds the {self.max_bytes // (1024 * 1024)} MB limit.",
                code="document_too_large",
            )
        suffix = Path(safe_name).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise CareerDomainError(
                "Unsupported document type. Use TXT, Markdown, PDF, or DOCX.",
                code="unsupported_document_type",
            )
        if suffix in {".txt", ".md", ".markdown"}:
            text = self._decode_text(content)
            parser_name = "plain_text_v1"
        elif suffix == ".pdf":
            text = self._parse_pdf(content)
            parser_name = "pypdf_v1"
        else:
            text = self._parse_docx(content)
            parser_name = "python_docx_v1"
        normalized = self._normalize(text)
        if not normalized:
            raise CareerDomainError(
                "No readable text was found in the document.", code="document_has_no_text"
            )
        return ParsedDocument(
            file_name=safe_name,
            media_type=media_type or self._media_type(suffix),
            text=normalized,
            parser_name=parser_name,
        )

    def parse_pasted_text(self, *, name: str, text: str) -> ParsedDocument:
        safe_name = self._validate_file_name(name if Path(name).suffix else f"{name}.txt")
        encoded = text.encode("utf-8")
        if len(encoded) > self.max_bytes:
            raise CareerDomainError("Pasted text is too large.", code="document_too_large")
        normalized = self._normalize(text)
        if not normalized:
            raise CareerDomainError("Pasted text is empty.", code="empty_document")
        return ParsedDocument(safe_name, "text/plain", normalized, "pasted_text_v1")

    @staticmethod
    def _validate_file_name(file_name: str) -> str:
        candidate = file_name.strip()
        if not candidate or candidate in {".", ".."}:
            raise CareerDomainError("A valid file name is required.", code="invalid_file_name")
        if "/" in candidate or "\\" in candidate or Path(candidate).name != candidate:
            raise CareerDomainError(
                "File paths are not accepted as file names.", code="invalid_file_name"
            )
        cleaned = re.sub(r"[\x00-\x1f<>:\"|?*]", "_", candidate).strip(" .")
        if not cleaned:
            raise CareerDomainError("A valid file name is required.", code="invalid_file_name")
        return cleaned[:255]

    @staticmethod
    def _decode_text(content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                pass
        detected = chardet.detect(content)
        encoding = detected.get("encoding")
        confidence = float(detected.get("confidence") or 0)
        if not encoding or confidence < 0.7:
            raise CareerDomainError(
                "Text encoding could not be identified safely. Convert the file to UTF-8.",
                code="unsupported_text_encoding",
            )
        try:
            return content.decode(encoding)
        except (LookupError, UnicodeDecodeError) as exc:
            raise CareerDomainError(
                "Text encoding could not be decoded.", code="unsupported_text_encoding"
            ) from exc

    @staticmethod
    def _parse_pdf(content: bytes) -> str:
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise CareerDomainError(
                    "Encrypted PDF documents are not supported.", code="encrypted_document"
                )
            if len(reader.pages) > 100:
                raise CareerDomainError("PDF has too many pages.", code="document_too_large")
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except CareerDomainError:
            raise
        except Exception as exc:
            raise CareerDomainError(
                "The PDF could not be parsed.", code="document_parse_failed"
            ) from exc

    @staticmethod
    def _parse_docx(content: bytes) -> str:
        try:
            document = DocxDocument(BytesIO(content))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        except Exception as exc:
            raise CareerDomainError(
                "The DOCX could not be parsed.", code="document_parse_failed"
            ) from exc

    @staticmethod
    def _normalize(text: str) -> str:
        text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
        return "\n".join(lines).strip()

    @staticmethod
    def _media_type(suffix: str) -> str:
        return {
            ".txt": "text/plain",
            ".md": "text/markdown",
            ".markdown": "text/markdown",
            ".pdf": "application/pdf",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }[suffix]
